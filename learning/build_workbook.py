"""Build the study workbook from its editable Markdown source using python-docx."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(0.75)
sec.left_margin = sec.right_margin = Inches(0.8)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.13
normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
for name in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.name = 'Calibri'
    doc.styles[name].font.color.rgb = RGBColor.from_string('123B52')
doc.styles['Heading 1'].font.size = Pt(21)
doc.styles['Heading 1'].paragraph_format.page_break_before = True
doc.styles['Heading 2'].font.size = Pt(14)
code_style = doc.styles.add_style('Code example', 1)
code_style.font.name = 'Consolas'
code_style.font.size = Pt(8.5)
code_style.paragraph_format.space_after = Pt(0)
code_style.paragraph_format.line_spacing = 1.0
code_style.paragraph_format.left_indent = Inches(0.12)
footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('DataWorks study workbook  |  ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.core_properties.title = 'Reservoir Data Science: A Workbook from Zero'
doc.core_properties.subject = 'T1–T4, Power BI and technical storytelling'
doc.core_properties.author = 'Choongmh portfolio study edition'

def paragraph(text, style=None):
    p=doc.add_paragraph(style=style)
    for i,part in enumerate(re.split(r'`([^`]+)`',text)):
        r=p.add_run(part)
        if i%2:
            r.font.name='Consolas'; r.font.size=Pt(9.5)
    return p

text=(HERE/'workbook.md').read_text(encoding='utf-8')
lines=text.splitlines()
doc.add_heading('Reservoir Data Science',0)
doc.add_paragraph('A workbook from zero',style='Subtitle')
doc.add_paragraph('DataWorks Challenge 2026 | Personal study edition for Choongmh')
doc.add_paragraph('T1 data preparation · T2 machine learning · T3 forecasting · T4 Power BI and storytelling')
doc.add_paragraph('English explanations with Chinese support for technical vocabulary. Worked calculations, practical exercises, answer explanations and a complete project reading path.')
doc.add_paragraph('21 October 2026')
doc.add_heading('Learning route',1)
for line in lines:
    if line.startswith('# ') and re.match(r'# \d+\.',line):
        paragraph(line[2:])
in_code=False
for line in lines[1:]:
    if line.startswith('```'):
        in_code=not in_code
        if not in_code: doc.add_paragraph().paragraph_format.space_after=Pt(1)
        continue
    if in_code:
        p=doc.add_paragraph(line,style='Code example')
        shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'F1F5F8'); p._p.get_or_add_pPr().append(shade)
    elif line.startswith('# '): doc.add_heading(line[2:],1)
    elif line.startswith('## '): doc.add_heading(line[3:],2)
    elif line.strip(): paragraph(line)

doc.add_heading('Visual reading practice',1)
figures=[
 ('baseline_top10.png','Baseline history matches','Find the observed line and the simulator lines. Identify which stream shows the largest visible mismatch. Use the original axis units.'),
 ('proxy_test_rates.png','Proxy rate diagnostics','Compare proxy and simulator rates. Explain why this tests historical emulation rather than future production accuracy.'),
 ('t3_forecast_all.png','Future scenario curves','Find the boundary between history and extrapolation. Explain which assumptions create the future curves.'),
 ('t3_water_cap_sensitivity.png','Water-cap sensitivity','Compare the effect of cap choices. Explain why this spread is an assumption comparison, not a calibrated confidence interval.')
]
for filename,title,task in figures:
    path=HERE.parent/'outputs/figures'/filename
    if path.exists():
        doc.add_heading(title,2)
        paragraph(task)
        doc.add_picture(str(path),width=Inches(6.4))
        p=paragraph('Source: outputs/figures/'+filename)
        p.style='Caption'
out=HERE/'DWC2026_Learning_Workbook.docx'
doc.save(out)
print('Saved',out)
print('Words in learning source:',len(text.split()))
print('Paragraphs:',len(doc.paragraphs),'Figures:',len(doc.inline_shapes))
