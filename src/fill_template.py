"""T3 Task G Part 2: fill the 10-sheet deliverable template.

Reads:  outputs/t3_forecast_all.csv  (scenario 'main', ids 1-10)
Reads:  data/raw/09 Template Deliverable.xlsx  (structure only; original never modified)
Writes: outputs/09 Template Deliverable_filled.xlsx
Writes: outputs/t3_template_mapping.csv

Sheet mapping: id i  ->  sheet 'Case i'.

Column layout (A-D; column E in Cases 1-3 is left untouched):
  A: Date [text "mm/dd/yyyy"]
  B: Gas  production cumulative [MSCF]
  C: Oil  production cumulative [STB]
  D: Water production cumulative [STB]

Case 1 special: template has only 40 data rows (2008-04-01 to 2018-01-01).
                40 rows are appended to reach the full 80 (2008-04-01 to 2028-01-01).

Run: .venv/bin/python -m src.fill_template
"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import pandas as pd

from src import config as C

_TEMPLATE    = C.DATA_RAW   / "09 Template Deliverable.xlsx"
_OUT_FILLED  = C.OUTPUTS_DIR / "09 Template Deliverable_filled.xlsx"
_MAPPING_CSV = C.OUTPUTS_DIR / "t3_template_mapping.csv"

_FC_START = pd.Timestamp("2008-04-01")
_FC_END   = pd.Timestamp("2028-01-01")

# Columns with E empty in Cases 1-3 (max_col=5) and absent in Cases 4-10 (max_col=4).
_COL_DATE  = 1   # A
_COL_GAS   = 2   # B  Gas cum [MSCF]
_COL_OIL   = 3   # C  Oil cum [STB]
_COL_WATER = 4   # D  Water cum [STB]


def _fmt_date(d: pd.Timestamp) -> str:
    return d.strftime("%m/%d/%Y")


def main() -> None:  # noqa: C901
    fc = pd.read_csv(C.OUTPUTS_DIR / "t3_forecast_all.csv", parse_dates=["date"])
    main_fc = fc[(fc["scenario"] == "main") &
                 (fc["date"] > pd.Timestamp("2008-01-01"))].copy()
    ids = sorted(main_fc["id"].unique())
    assert ids == list(range(1, 11)), f"Expected ids 1-10, got {ids}"

    fc_dates = pd.date_range(_FC_START, _FC_END, freq="QS")
    assert len(fc_dates) == 80

    # Copy template; fill the copy (never touch the original)
    shutil.copy2(_TEMPLATE, _OUT_FILLED)
    wb = openpyxl.load_workbook(_OUT_FILLED)

    mapping_rows = []

    for id_ in ids:
        ws = wb[f"Case {id_}"]
        id_data = main_fc[main_fc["id"] == id_].set_index("date").sort_index()
        assert len(id_data) == 80, f"id={id_}: {len(id_data)} rows (expected 80)"

        for row_idx, date in enumerate(fc_dates, start=2):
            row   = id_data.loc[date]
            gas_v = float(row["gas_cum"])
            oil_v = float(row["oil_cum"])
            wat_v = float(row["water_cum"])
            dt_tx = _fmt_date(date)

            ws.cell(row=row_idx, column=_COL_DATE).value  = dt_tx
            ws.cell(row=row_idx, column=_COL_GAS).value   = gas_v
            ws.cell(row=row_idx, column=_COL_OIL).value   = oil_v
            ws.cell(row=row_idx, column=_COL_WATER).value = wat_v
            # Column E (5) is never written -- left as-is.

            mapping_rows.append({
                "id":            id_,
                "sheet":         f"Case {id_}",
                "excel_row":     row_idx,
                "date":          date.strftime("%Y-%m-%d"),
                "gas_cum_MSCF":  round(gas_v, 0),
                "oil_cum_STB":   round(oil_v, 0),
                "water_cum_STB": round(wat_v, 0),
            })

    wb.save(_OUT_FILLED)
    wb.close()
    print(f"Saved {_OUT_FILLED}", flush=True)

    mapping_df = pd.DataFrame(mapping_rows)
    mapping_df.to_csv(_MAPPING_CSV, index=False)
    print(f"Saved mapping ({len(mapping_df):,} rows) -> {_MAPPING_CSV}", flush=True)

    _verify(_OUT_FILLED, main_fc, fc_dates)


def _verify(path: Path, main_fc: pd.DataFrame,
            fc_dates: pd.DatetimeIndex) -> None:
    print("\n=== PASS/FAIL checks ===")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    all_ok = True

    for id_ in range(1, 11):
        sname = f"Case {id_}"
        ws    = wb[sname]
        notes = []

        # 1. Row count (header + 80 data = 81)
        if ws.max_row != 81:
            notes.append(f"row_count={ws.max_row} (want 81)")

        # 2. First date = "04/01/2008"
        first_d = ws.cell(row=2, column=_COL_DATE).value
        if first_d != "04/01/2008":
            notes.append(f"first_date='{first_d}'")

        # 3. Last date = "01/01/2028"
        last_d = ws.cell(row=81, column=_COL_DATE).value
        if last_d != "01/01/2028":
            notes.append(f"last_date='{last_d}'")

        # 4. All date cells are text strings in "mm/dd/yyyy" format
        bad_date = None
        for r in range(2, 82):
            v = ws.cell(row=r, column=_COL_DATE).value
            if not isinstance(v, str):
                bad_date = (r, v); break
            try:
                pd.to_datetime(v, format="%m/%d/%Y")
            except Exception:
                bad_date = (r, v); break
        if bad_date:
            notes.append(f"bad_date_cell row={bad_date[0]} val='{bad_date[1]}'")

        # 5. Values non-null and non-negative
        bad_val = None
        for r in range(2, 82):
            for col in (_COL_GAS, _COL_OIL, _COL_WATER):
                v = ws.cell(row=r, column=col).value
                if v is None or float(v) < 0:
                    bad_val = (r, col, v); break
            if bad_val:
                break
        if bad_val:
            notes.append(f"bad_value row={bad_val[0]} col={bad_val[1]} val={bad_val[2]}")

        # 6. Column E untouched (None) in Cases 1-3
        if id_ <= 3:
            bad_e = next(
                (r for r in range(2, 82)
                 if ws.cell(row=r, column=5).value is not None), None)
            if bad_e is not None:
                notes.append(f"col_E_modified at row={bad_e}")

        # 7. Spot check: last row values match forecast CSV
        id_data = main_fc[main_fc["id"] == id_].set_index("date")
        ref     = id_data.loc[pd.Timestamp("2028-01-01")]
        tol     = 1.0
        for col_idx, col_name in [(_COL_GAS, "gas_cum"),
                                   (_COL_OIL, "oil_cum"),
                                   (_COL_WATER, "water_cum")]:
            cell_v = float(ws.cell(row=81, column=col_idx).value)
            ref_v  = float(ref[col_name])
            if abs(cell_v - ref_v) > tol:
                notes.append(f"value_mismatch {col_name} cell={cell_v:.0f} ref={ref_v:.0f}")

        status = "PASS" if not notes else "FAIL: " + "; ".join(notes)
        e_note = " (col E untouched)" if id_ <= 3 and not notes else ""
        print(f"  {sname}: {status}{e_note}")
        if notes:
            all_ok = False

    wb.close()
    print(f"\n{'ALL PASS' if all_ok else 'SOME CHECKS FAILED'}")


if __name__ == "__main__":
    main()
