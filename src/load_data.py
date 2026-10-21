"""
Step 0: Data preparation (Track 1).

Reads the organiser Excel files ONCE and writes clean parquet files to data/processed/.
Everyone else reads the parquet files through the `read_*` helpers at the bottom.

Usage (from the repo root):
    python -m src.load_data            # params, curves, history, master table  (~1 min)
    python -m src.load_data --grid     # ALSO convert the 3D grid files          (~40 min, optional)

Outputs (data/processed/):
    params.parquet         100 rows: case, split, 4 parameters                         -> X
    curves_long.parquet    100 cases x 41 dates: 6 streams (cum + rate)               -> y (long form)
    history.parquet        41 dates: REAL field data, 6 streams + water injection     -> answer key
    master_table.parquet   100 rows: case, split, 4 params, 123 y columns (cum)       -> T2 training table
    master_table.csv       same as above, for quick viewing
    grid/case_###.parquet  (optional) one 3D grid per case, -999.25 replaced by NaN
"""
import argparse
import re
import warnings

import numpy as np
import openpyxl
import pandas as pd

from src import config as C


# ============================================================ 1. parameters (X)
def load_params() -> pd.DataFrame:
    """100 cases x 4 uncertainty parameters, plus case number and split."""
    df = pd.read_excel(C.raw_file(C.INTRO_PREFIX), sheet_name="Uncertainty Values")
    df["case"] = df["Item"].str.extract(r"(\d+)").astype(int)
    df = df.rename(columns=C.PARAMS)
    df["split"] = df["case"].map(C.split_of)
    df = df[["case", "split"] + C.PARAM_COLS].sort_values("case").reset_index(drop=True)

    assert len(df) == C.N_CASES, f"Expected {C.N_CASES} cases, got {len(df)}"
    assert df[C.PARAM_COLS].notna().all().all(), "Missing parameter values"
    return df


# ============================================================ 2. production curves (y) + history
def _read_curve_sheet(sheet: str) -> pd.DataFrame:
    """Read one curve sheet and return a table: rows = the 41 history dates,
    columns = Case_1..Case_100 + RAMP_History.

    Sheet layout: row 1 = codes (X, Y1..), row 2 = names (Date, Case_1..), row 3 = units, data from row 4.
    Each case only has values on ITS OWN simulator dates (blank elsewhere), so every column is
    handled separately: drop its blanks, then take the value on each history date.
    """
    raw = pd.read_excel(C.raw_file(C.PRODUCTION_PREFIX), sheet_name=sheet, header=1)
    raw = raw.dropna(subset=["Date"])                       # drops the units row
    raw["Date"] = pd.to_datetime(raw["Date"])
    raw = raw.set_index("Date").apply(pd.to_numeric, errors="coerce")

    target = C.HISTORY_DATES
    out = {}
    for col in raw.columns:
        s = raw[col].dropna()
        s = s[~s.index.duplicated(keep="last")]
        aligned = s.reindex(target)                         # exact lookup on the 41 dates
        if aligned.isna().any():                            # safety net: should not happen with this data
            missing = aligned.index[aligned.isna()]
            warnings.warn(f"[{sheet}] {col}: {len(missing)} history date(s) missing, filled by interpolation")
            x_known = s.index.values.astype("int64")
            aligned = pd.Series(np.interp(target.values.astype("int64"), x_known, s.values), index=target)
        out[col] = aligned
    return pd.DataFrame(out, index=target)


def add_rates(df: pd.DataFrame) -> pd.DataFrame:
    """Add oil/gas/water rates derived from the cumulative columns (see config.RATE_STREAMS note):
    rate(t) = (cum(t) - cum(t-1)) / days in that quarter. The first date gets 0."""
    days = pd.Series(C.HISTORY_DATES).diff().dt.days.to_numpy()   # NaN, 90, 91, 92, ...
    for cum, rate in zip(C.CUM_STREAMS, C.RATE_STREAMS):
        vals = df[cum].to_numpy().reshape(-1, len(C.HISTORY_DATES))  # (n_series, 41)
        r = np.diff(vals, axis=1, prepend=vals[:, :1]) / np.nan_to_num(days, nan=1.0)
        df[rate] = r.reshape(-1)
    return df


def load_curves_and_history() -> tuple[pd.DataFrame, pd.DataFrame]:
    """Returns
    curves_long : case, split, date, oil_cum, gas_cum, water_cum, oil_rate, gas_rate, water_rate
    history     : date, the same 6 streams (REAL data) + the rates exactly as reported in the
                  'Field Production' sheet (*_reported, quarter STARTING at the date) + water injection
    """
    case_frames, hist = [], {}
    for stream, sheet in C.CUM_SHEETS.items():
        wide = _read_curve_sheet(sheet)
        hist[stream] = wide.pop(C.HISTORY_COL)
        wide.columns = wide.columns.str.extract(r"(\d+)", expand=False).astype(int)
        long = wide.stack().rename(stream)
        long.index = long.index.set_names(["date", "case"])
        case_frames.append(long)

    curves = pd.concat(case_frames, axis=1).reset_index()
    curves = curves.sort_values(["case", "date"]).reset_index(drop=True)
    curves = add_rates(curves)
    curves["split"] = curves["case"].map(C.split_of)
    curves = curves[["case", "split", "date"] + C.STREAMS]

    history = add_rates(pd.DataFrame(hist).rename_axis("date").reset_index())

    # Reported rates + water injection come from the "Field Production" sheet (forward convention).
    fp = pd.read_excel(C.raw_file(C.PRODUCTION_PREFIX), sheet_name=C.HISTORY_SHEET)
    fp["Date"] = pd.to_datetime(fp["Date"])
    fp = fp.set_index("Date").reindex(C.HISTORY_DATES)
    reported = {"oil_rate": "Oil production rate (STB/d)",
                "gas_rate": "Gas production rate (MSCF/d)",
                "water_rate": "Water production rate (STB/d)"}
    for stream, col in reported.items():
        history[f"{stream}_reported"] = fp[col].to_numpy()
        # check: our derived rate at t must equal the reported rate one quarter earlier
        derived, rep_prev = history[stream].to_numpy()[1:], fp[col].to_numpy()[:-1]
        if not np.allclose(derived, rep_prev, rtol=1e-3, atol=1.0):
            warnings.warn(f"Derived {stream} does not match the reported rates shifted by one quarter")
    history["water_injection_rate_reported"] = fp["Water injection rate (STB/d)"].to_numpy()
    return curves, history


# ============================================================ 3. master table (one row per case)
def y_columns(streams=C.CUM_STREAMS) -> list[str]:
    """Names of the y columns in the master table, e.g. oil_cum_t01 ... water_cum_t41."""
    n = len(C.HISTORY_DATES)
    return [f"{s}_t{i:02d}" for s in streams for i in range(1, n + 1)]


def build_master_table(params: pd.DataFrame, curves: pd.DataFrame) -> pd.DataFrame:
    """case | split | 4 params || oil_cum_t01..t41 | gas_cum_t01..t41 | water_cum_t01..t41
    t01 = 1998-01-01 ... t41 = 2008-01-01 (see config.HISTORY_DATES)."""
    date_idx = {d: i + 1 for i, d in enumerate(C.HISTORY_DATES)}
    long = curves.melt(id_vars=["case", "date"], value_vars=C.CUM_STREAMS, var_name="stream")
    long["col"] = long["stream"] + "_t" + long["date"].map(date_idx).map("{:02d}".format)
    wide = long.pivot(index="case", columns="col", values="value")[y_columns()]
    master = params.merge(wide, left_on="case", right_index=True)
    assert master.shape == (C.N_CASES, 2 + len(C.PARAM_COLS) + len(y_columns()))
    return master


# ============================================================ 4. checks
def sanity_checks(params, curves, history, master) -> None:
    assert curves[C.STREAMS].notna().all().all(), "NaN in curves"
    assert history.drop(columns="date").notna().all().all(), "NaN in history"
    assert len(curves) == C.N_CASES * len(C.HISTORY_DATES)
    for s in C.CUM_STREAMS:   # cumulative values must never go down
        dec = curves.groupby("case")[s].apply(lambda v: (np.diff(v) < -1e-6).any())
        assert not dec.any(), f"{s} decreases for cases {list(dec[dec].index)}"
    counts = params["split"].value_counts().to_dict()
    assert counts == {"train": 70, "validation": 15, "test": 15}, counts

    print("\n=== Data summary ===")
    print(f"params        {params.shape}   split {counts}")
    print(f"curves_long   {curves.shape}   dates {curves.date.min():%Y-%m-%d} -> {curves.date.max():%Y-%m-%d}")
    print(f"history       {history.shape}")
    print(f"master_table  {master.shape}   (2 id + {len(C.PARAM_COLS)} X + {len(y_columns())} y)")
    end = curves[curves.date == C.HISTORY_DATES[-1]]
    h_end = history.iloc[-1]
    print("\nCumulative at 2008-01-01: cases min / max  vs  REAL")
    for s in C.CUM_STREAMS:
        print(f"  {s:10s} {end[s].min():>14,.0f} / {end[s].max():>14,.0f}   vs {h_end[s]:>14,.0f}")
    print("\nParameter ranges (allowed search space for optimisation):")
    print(params[C.PARAM_COLS].agg(["min", "max"]).T.to_string())


# ============================================================ 5. grid files (optional, slow)
def _short_name(col: str) -> str:
    """'Porosity (PORO)' -> 'PORO';  'I Index' -> 'I';  'X Coordinate' -> 'X'."""
    m = re.search(r"\(([A-Z]+)\)", col)
    return m.group(1) if m else col.strip().split()[0]


def convert_grid_files(overwrite: bool = False) -> None:
    """Convert every case sheet of files 03-08 to grid/case_###.parquet (resumable: skips existing)."""
    C.GRID_DIR.mkdir(parents=True, exist_ok=True)
    for prefix in C.GRID_PREFIXES:
        try:
            path = C.raw_file(prefix)
        except FileNotFoundError as e:
            warnings.warn(str(e))
            continue
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in wb.sheetnames:
            case = int(re.search(r"\d+", sheet).group())
            out = C.GRID_DIR / f"case_{case:03d}.parquet"
            if out.exists() and not overwrite:
                continue
            rows = wb[sheet].iter_rows(values_only=True)
            header = [_short_name(h) for h in next(rows)]
            df = pd.DataFrame(rows, columns=header, dtype="float64").replace(C.GRID_NA, np.nan)
            df[["I", "J", "K"]] = df[["I", "J", "K"]].astype("int16")
            df.insert(0, "case", case)
            df.to_parquet(out, index=False)
            print(f"  {path.name} / {sheet} -> {out.name}  ({C.split_of(case)}, {len(df):,} cells)")
        wb.close()


# ============================================================ 6. helpers for the team
def read_params() -> pd.DataFrame:
    return pd.read_parquet(C.DATA_PROCESSED / "params.parquet")


def read_curves() -> pd.DataFrame:
    return pd.read_parquet(C.DATA_PROCESSED / "curves_long.parquet")


def read_history() -> pd.DataFrame:
    return pd.read_parquet(C.DATA_PROCESSED / "history.parquet")


def read_master() -> pd.DataFrame:
    return pd.read_parquet(C.DATA_PROCESSED / "master_table.parquet")


def read_grid(case: int) -> pd.DataFrame:
    return pd.read_parquet(C.GRID_DIR / f"case_{case:03d}.parquet")


# ============================================================ main
def main(grid: bool = False) -> None:
    C.DATA_PROCESSED.mkdir(parents=True, exist_ok=True)
    print("Reading parameters ...")
    params = load_params()
    print("Reading production curves (3 cumulative sheets; rates derived) ...")
    curves, history = load_curves_and_history()
    master = build_master_table(params, curves)
    sanity_checks(params, curves, history, master)

    params.to_parquet(C.DATA_PROCESSED / "params.parquet", index=False)
    curves.to_parquet(C.DATA_PROCESSED / "curves_long.parquet", index=False)
    history.to_parquet(C.DATA_PROCESSED / "history.parquet", index=False)
    master.to_parquet(C.DATA_PROCESSED / "master_table.parquet", index=False)
    master.to_csv(C.DATA_PROCESSED / "master_table.csv", index=False)
    print(f"\nSaved to {C.DATA_PROCESSED}")

    if grid:
        print("\nConverting grid files (slow, resumable) ...")
        convert_grid_files()


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="DWC2026 data preparation")
    ap.add_argument("--grid", action="store_true", help="also convert the 3D grid files (~40 min)")
    main(grid=ap.parse_args().grid)